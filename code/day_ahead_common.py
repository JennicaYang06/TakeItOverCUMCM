"""
问题2 公共基础设施：数据路径、日前 MILP（问题1结构+全年滚动衔接）、报童安全边际、
result2.xlsx 模板导出、论文素材汇总、代表日画图。

这些逻辑与"用什么方法预测负荷/光伏"无关，被多个问题2版本（如 problem2.py 的
Holt-Winters+晴空包络版、forecast_lightgbm_最终版.py 的 LightGBM 版）共用，
确保除预测方法外其余建模假设完全一致，比较才有意义。
"""
from pathlib import Path
import datetime as dt
import numpy as np
import pandas as pd
import cvxpy as cp

# ===================== 公共参数（附录 1）=====================
DT = 1.0 / 6.0
SOC_MIN, SOC_MAX = 1200.0, 10800.0
SOC0_INIT = 6000.0
P_MAX = 5000.0
E_STEP_MAX = P_MAX * DT
RT_EFF = 0.90
ETA_C = np.sqrt(RT_EFF)
ETA_D = np.sqrt(RT_EFF)
T = 144
EMERGENCY_MULT = 5.0
SOLVER = "HIGHS"          # 本机没装 GLPK_MI，用开源 HIGHS（MILP 求解快且稳定）

# 报童临界分位数：欠量单位成本 Cu=(EMERGENCY_MULT-1)*p，超量单位成本 Co=p，
# q*=Cu/(Cu+Co)，与价格 p 本身无关。
SAFETY_QUANTILE = (EMERGENCY_MULT - 1) / EMERGENCY_MULT   # = 0.8
SAFETY_WINDOW = 60

EXPORT_START = "2025-02-01"
EXPORT_END = "2025-12-31"
REP_DATES = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]

TBL1 = [("10:00-10:10", 60), ("12:00-12:10", 72), ("14:00-14:10", 84),
        ("16:00-16:10", 96), ("18:00-18:10", 108), ("20:00-20:10", 120)]
TBL2 = [("0:00-4:00", 0, 24), ("4:00-8:00", 24, 48), ("8:00-12:00", 48, 72),
        ("12:00-16:00", 72, 96), ("16:00-20:00", 96, 120), ("20:00-24:00", 120, 144)]


# ======================= 数据路径 =======================
def find_dir(candidates):
    for c in candidates:
        if c.exists():
            return c
    raise FileNotFoundError("找不到数据目录，尝试过：\n" + "\n".join(str(c) for c in candidates))


def resolve_data_dirs(here: Path):
    """here: 调用脚本所在目录（code/）。返回 (root, data_dir, template_dir)。"""
    root = here.parent
    data_dir = find_dir([here / "附件", root / "题目和附件" / "C题" / "附件"])
    template_dir = find_dir([data_dir / "附件5", root / "题目和附件" / "C题" / "附件" / "附件5"])
    return root, data_dir, template_dir


# ======================= 基础工具函数 =======================
def adjacent_average(arr):
    """首尾相接平均：把整点瞬时功率折算成"区间平均功率"，数据点数不变。
    用于完整一天(144点)的循环数据：第1格用"昨天最后一点"(即arr[-1])做左端点。"""
    arr = np.asarray(arr, dtype=float)
    result = np.zeros(len(arr))
    result[0] = (arr[-1] + arr[0]) / 2
    result[1:] = (arr[:-1] + arr[1:]) / 2
    return result


def adjacent_average_with_prev(arr, prev_value):
    """同 adjacent_average，但左端点显式传入（用于非整天的分段序列，不能用 arr[-1] 兜底）。"""
    arr = np.asarray(arr, dtype=float)
    result = np.zeros(len(arr))
    result[0] = (prev_value + arr[0]) / 2
    result[1:] = (arr[:-1] + arr[1:]) / 2
    return result


def interp_hourly_to_10min(anchor_value, hourly_values, n_blocks):
    """把"整点瞬时功率预报"线性插值到10分钟粒度的瞬时功率点序列。
    anchor_value: 区间起点(当前时刻)的已知瞬时功率。
    hourly_values: 长度 n_blocks//6 的整点预报（第1个=1小时后，第2个=2小时后...）。
    返回长度 n_blocks 的瞬时功率序列，对应 10,20,...,n_blocks*10 分钟处的取值
    （与附件1/2原始数据"每列代表某整10分钟时刻瞬时功率"的约定一致）。"""
    n_hours = len(hourly_values)
    assert n_blocks == n_hours * 6, f"n_blocks={n_blocks} 应等于 hourly_values长度({n_hours})*6"
    anchor_minutes = np.arange(0, n_hours + 1) * 60.0
    anchor_vals = np.concatenate([[anchor_value], np.asarray(hourly_values, dtype=float)])
    sample_minutes = np.arange(1, n_blocks + 1) * 10.0
    return np.interp(sample_minutes, anchor_minutes, anchor_vals)


def time_range_str(i):
    start_min, end_min = i * 10, (i + 1) * 10
    if end_min >= 1440:
        end_h, end_m = 24, 0
    else:
        end_h, end_m = end_min // 60, end_min % 60
    start_h, start_m = start_min // 60, start_min % 60
    return f"{start_h:02d}:{start_m:02d}-{end_h:02d}:{end_m:02d}"


def merge_deficit_segments(deficit_e, eps=1e-6):
    segs = []
    t = 0
    n = len(deficit_e)
    while t < n:
        if deficit_e[t] > eps:
            start = t
            s = 0.0
            while t < n and deficit_e[t] > eps:
                s += deficit_e[t]
                t += 1
            segs.append((start, t, s))
        else:
            t += 1
    return segs


def error_metrics(pred, actual, mape_thresh=None):
    err = np.asarray(pred, dtype=float) - np.asarray(actual, dtype=float)
    actual = np.asarray(actual, dtype=float)
    mae = float(np.mean(np.abs(err)))
    rmse = float(np.sqrt(np.mean(err ** 2)))
    mask = (actual > mape_thresh) if mape_thresh is not None else np.ones_like(actual, dtype=bool)
    mape = float(np.mean(np.abs(err[mask] / actual[mask])) * 100) if mask.sum() > 0 else float("nan")
    return mae, rmse, mape


# ======================= 读取附件1/附件2 =======================
def load_price_and_actuals(data_dir: Path):
    """返回电价（每天固定复用）与全年实际负荷/光伏。"""
    a1 = pd.read_excel(data_dir / "附件1.xlsx", header=None).values
    price_raw = a1[1:, 1].astype(float)
    load_kw_seed = a1[1:, 2].astype(float)
    pv_kw_seed = a1[1:, 3].astype(float)
    assert len(price_raw) == T, f"附件1 应为 144 行，实际 {len(price_raw)}"
    price_avg = adjacent_average(price_raw)

    xls2 = data_dir / "附件2.xlsx"
    df_load = pd.read_excel(xls2, sheet_name="小区负载", header=0)
    df_pv = pd.read_excel(xls2, sheet_name="光伏发电实际功率", header=0)
    dates = pd.to_datetime(df_load.iloc[:, 0]).dt.normalize().to_numpy()
    load_kw_all = df_load.iloc[:, 1:1 + T].to_numpy(dtype=float)
    pv_kw_all = df_pv.iloc[:, 1:1 + T].to_numpy(dtype=float)
    n_days = len(dates)
    assert load_kw_all.shape == (n_days, T) and pv_kw_all.shape == (n_days, T)
    dates_str = np.array([pd.Timestamp(d).strftime("%Y-%m-%d") for d in dates])

    return dict(price_raw=price_raw, price_avg=price_avg,
                load_kw_seed=load_kw_seed, pv_kw_seed=pv_kw_seed,
                dates_str=dates_str, load_kw_all=load_kw_all, pv_kw_all=pv_kw_all,
                n_days=n_days)


def compute_export_range(dates_str, n_days, export_start=EXPORT_START, export_end=EXPORT_END):
    export_start_idx = int(np.where(dates_str == export_start)[0][0])
    export_end_idx = int(np.where(dates_str == export_end)[0][0]) if export_end in dates_str else n_days - 1
    export_end_idx = min(export_end_idx, n_days - 1)
    do_export = export_start_idx < n_days
    return export_start_idx, export_end_idx, do_export


# ======================= 负荷预测（问题2、问题3共用：附件3只给光伏预报，负荷仍需自己预测） =======================
class LoadForecaster:
    """按144个时刻分别做加性 Holt-Winters（周期=7天，无趋势项）滚动预测，0:00做一次，全天不再修正。"""

    def __init__(self, seed_curve, alpha=0.25, gamma=0.25):
        self.alpha = alpha
        self.gamma = gamma
        self.level = np.asarray(seed_curve, dtype=float).copy()
        self.seasonal = np.zeros((len(seed_curve), 7))

    def predict(self, day_idx):
        dow = day_idx % 7
        return np.clip(self.level + self.seasonal[:, dow], 0, None)

    def update(self, day_idx, actual):
        dow = day_idx % 7
        pred = self.level + self.seasonal[:, dow]
        err = actual - pred
        self.level = self.level + self.alpha * err
        self.seasonal[:, dow] = self.seasonal[:, dow] + self.gamma * err


# ======================= 报童安全边际 =======================
class NetErrorTracker:
    """按144个时刻分别维护"净负荷(=负荷-光伏)预测残差"的滚动历史，供报童安全边际取分位数。"""

    def __init__(self, T=T, window=SAFETY_WINDOW, quantile=SAFETY_QUANTILE):
        self.window = window
        self.quantile = quantile
        self.buffer = []
        self.T = T

    def get_margin(self, day_idx=None):
        if not self.buffer:
            return np.zeros(self.T)
        arr = np.array(self.buffer[-self.window:])
        return np.quantile(arr, self.quantile, axis=0)

    def update(self, forecast_net_e, actual_net_e):
        self.buffer.append(actual_net_e - forecast_net_e)


# ======================= 每日 MILP（复用问题1结构） =======================
class DayAheadMILP:
    """
    单日 MILP：g(计划购电)/c(充电)/d(放电)/z(充放电互斥)/soc(储能电量)。
    soc0、负荷目标、光伏预测做成 cvxpy Parameter，编译一次、逐日复用求解，避免重复编译开销。
    与问题1的区别：去掉"soc[0]=soc[T]"的单日闭环，改为调用方自行传入前一天末电量作为 soc0。
    """

    def __init__(self, price_avg):
        self.price_avg = price_avg
        self.g = cp.Variable(T, nonneg=True, name="plan_buy")
        self.c = cp.Variable(T, nonneg=True, name="charge")
        self.d = cp.Variable(T, nonneg=True, name="discharge")
        self.z = cp.Variable(T, boolean=True, name="charge_discharge_flag")
        self.soc = cp.Variable(T + 1, name="soc")

        self.soc0_p = cp.Parameter(name="soc0")
        self.load_e_p = cp.Parameter(T, name="load_e")
        self.pv_e_p = cp.Parameter(T, name="pv_e")

        cons = [
            self.soc[0] == self.soc0_p,
            self.soc >= SOC_MIN,
            self.soc <= SOC_MAX,
            self.soc[1:] == self.soc[:-1] + ETA_C * self.c - self.d / ETA_D,
            self.c <= E_STEP_MAX,
            self.d <= E_STEP_MAX,
            self.g + self.pv_e_p + self.d - self.c >= self.load_e_p,
            self.c <= E_STEP_MAX * self.z,
            self.d <= E_STEP_MAX * (1 - self.z),
        ]
        self.prob = cp.Problem(cp.Minimize(price_avg @ self.g), cons)

    def solve(self, soc0_val, load_e_val, pv_e_val):
        self.soc0_p.value = float(soc0_val)
        self.load_e_p.value = load_e_val
        self.pv_e_p.value = pv_e_val
        self.prob.solve(solver=SOLVER)
        if self.prob.status not in ("optimal", "optimal_inaccurate"):
            raise RuntimeError(f"day solve failed: status={self.prob.status}")
        gv = np.clip(np.asarray(self.g.value).ravel(), 0.0, None)
        cv = np.clip(np.asarray(self.c.value).ravel(), 0.0, None)
        dv = np.clip(np.asarray(self.d.value).ravel(), 0.0, None)
        socv = np.asarray(self.soc.value).ravel()
        return gv, cv, dv, socv


# ======================= 问题3：日内调整 MILP（只重新决策"剩余时段"） =======================
class SegmentMILP:
    """
    问题3的日内调整：0:00的整天计划只是"计划购电量"（对外报告用，也是违约/超额费的基准）。
    6:00/12:00/18:00各拿到一份新的光伏预报后，只对"接下来这一段"（默认6小时=36格）重新
    决策 g_adj/c/d/soc，负荷预测维持0点做的那版不变（附件3只更新光伏，没有给负荷的日内预报）。

    费用只算"调整相对计划"的差额：调整量比计划多的部分按1.5倍价格多付，
    调整量比计划少的部分按0.5倍价格计入违约金（即只按0.5倍价格"补偿"少买的部分，
    相当于比完全不调整省下0.5倍价格）。计划购电本身的费用(price*g_plan)由调用方另外累加，
    不放在这个子问题的目标函数里（对本段决策变量而言是常数，不影响最优解)。

    g_adj 与 (up,down) 的关系用等式约束 g_adj-g_plan_orig_p == up-down 显式表达，
    不能写成 cp.pos(g_adj-g_plan_orig)-cp.pos(g_plan_orig-g_adj) 直接相减——那样一凸一凹，
    不满足DCP。这里两个辅助变量的目标系数一正一负、净系数为正(1.5-0.5=1>0)，
    最优解处仍会自动收敛到 up=max(Δ,0)、down=max(-Δ,0) 的精确分解（详见问题3代码顶部注释推导）。
    """

    def __init__(self, seg_len):
        self.seg_len = seg_len
        n = seg_len
        self.g_adj = cp.Variable(n, nonneg=True, name="g_adj")
        self.up = cp.Variable(n, nonneg=True, name="adj_up")
        self.down = cp.Variable(n, nonneg=True, name="adj_down")
        self.c = cp.Variable(n, nonneg=True, name="charge")
        self.d = cp.Variable(n, nonneg=True, name="discharge")
        self.z = cp.Variable(n, boolean=True, name="charge_discharge_flag")
        self.soc = cp.Variable(n + 1, name="soc")

        self.soc0_p = cp.Parameter(name="soc0")
        self.load_e_p = cp.Parameter(n, name="load_e")
        self.pv_e_p = cp.Parameter(n, name="pv_e")
        self.g_plan_p = cp.Parameter(n, name="g_plan_orig")
        self.price_p = cp.Parameter(n, nonneg=True, name="price_seg")

        cons = [
            self.soc[0] == self.soc0_p,
            self.soc >= SOC_MIN,
            self.soc <= SOC_MAX,
            self.soc[1:] == self.soc[:-1] + ETA_C * self.c - self.d / ETA_D,
            self.c <= E_STEP_MAX,
            self.d <= E_STEP_MAX,
            self.g_adj + self.pv_e_p + self.d - self.c >= self.load_e_p,
            self.c <= E_STEP_MAX * self.z,
            self.d <= E_STEP_MAX * (1 - self.z),
            self.g_adj - self.g_plan_p == self.up - self.down,
        ]
        adj_cost = self.price_p @ (1.5 * self.up - 0.5 * self.down)
        self.prob = cp.Problem(cp.Minimize(adj_cost), cons)

    def solve(self, soc0_val, load_e_val, pv_e_val, g_plan_orig_val, price_seg_val):
        self.soc0_p.value = float(soc0_val)
        self.load_e_p.value = load_e_val
        self.pv_e_p.value = pv_e_val
        self.g_plan_p.value = g_plan_orig_val
        self.price_p.value = price_seg_val
        self.prob.solve(solver=SOLVER)
        if self.prob.status not in ("optimal", "optimal_inaccurate"):
            raise RuntimeError(f"segment solve failed: status={self.prob.status}")
        g_adj = np.clip(np.asarray(self.g_adj.value).ravel(), 0.0, None)
        cv = np.clip(np.asarray(self.c.value).ravel(), 0.0, None)
        dv = np.clip(np.asarray(self.d.value).ravel(), 0.0, None)
        socv = np.asarray(self.soc.value).ravel()
        adj_cost = float(self.prob.value)
        return g_adj, cv, dv, socv, adj_cost


# ======================= 结算一天（计划执行 + 实际值回代 + 紧急购电） =======================
def settle_day(gv, cv, dv, actual_load_e, actual_pv_e, price_avg):
    actual_supply = gv + actual_pv_e + dv - cv
    deficit_e = np.clip(actual_load_e - actual_supply, 0.0, None)
    emerg_cost = float((EMERGENCY_MULT * price_avg * deficit_e).sum())
    plan_cost = float(price_avg @ gv)
    return deficit_e, plan_cost, emerg_cost


# ======================= 导出 result2.xlsx / result3.xlsx 模板（共用子函数） =======================
def _fill_purchase_sheet(ws, results, key, export_start_idx, export_end_idx, price_avg, n_days):
    """填一张"购电量"格式的sheet（144格+下一天首格+全天购电量+全天购电费）。
    key: results[i] 中购电量数组的字段名（如 'gv'、'g_plan'、'g_final'）。"""
    for i in range(export_start_idx, export_end_idx + 1):
        r = results[i]
        row = 2 + (i - export_start_idx)
        gv = r[key]
        for col in range(2, 145):          # col2..144 -> gv[1..143]
            ws.cell(row=row, column=col, value=round(float(gv[col - 1]), 6))
        next_gv0 = results[i + 1][key][0] if i + 1 < n_days else gv[0]
        ws.cell(row=row, column=145, value=round(float(next_gv0), 6))
        ws.cell(row=row, column=146, value=round(float(gv.sum()), 6))
        ws.cell(row=row, column=147, value=round(float(price_avg @ gv), 6))


def _fill_charge_sheet(ws, results, export_start_idx, export_end_idx, c_key="cv", d_key="dv", soc_key="socv"):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    for i in range(export_start_idx, export_end_idx + 1):
        r = results[i]
        cv, dv, socv = r[c_key], r[d_key], r[soc_key]
        date_obj = dt.datetime.strptime(r["date"], "%Y-%m-%d")
        for bi, (label, a, b) in enumerate(TBL2):
            ws.cell(row=row, column=1, value=date_obj if bi == 0 else None)
            ws.cell(row=row, column=2, value=label)
            ws.cell(row=row, column=3, value=round(float(cv[a:b].sum()), 6))
            ws.cell(row=row, column=4, value=round(float(dv[a:b].sum()), 6))
            if bi == 0:
                ws.cell(row=row, column=5, value=dt.time(0, 0))
                ws.cell(row=row, column=6, value=round(float(socv[0]), 6))
            elif bi == 1:
                ws.cell(row=row, column=5, value="24:00")
                ws.cell(row=row, column=6, value=round(float(socv[-1]), 6))
            row += 1


def _fill_emergency_sheet(ws, results, export_start_idx, export_end_idx, deficit_key="deficit_e"):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    for i in range(export_start_idx, export_end_idx + 1):
        r = results[i]
        segs = merge_deficit_segments(r[deficit_key])
        if not segs:
            continue
        date_obj = dt.datetime.strptime(r["date"], "%Y-%m-%d")
        for si, (s, e, val) in enumerate(segs):
            ws.cell(row=row, column=1, value=date_obj if si == 0 else None)
            ws.cell(row=row, column=2, value=time_range_str(s)[:5] + "-" + time_range_str(e - 1)[6:])
            ws.cell(row=row, column=3, value=round(float(val), 6))
            row += 1


def export_result2_template(template_file, out_file, results, export_start_idx, export_end_idx,
                             price_avg, n_days):
    from openpyxl import load_workbook
    wb = load_workbook(template_file)
    _fill_purchase_sheet(wb["计划购电量"], results, "gv", export_start_idx, export_end_idx, price_avg, n_days)
    _fill_charge_sheet(wb["充放电量"], results, export_start_idx, export_end_idx)
    _fill_emergency_sheet(wb["紧急购电量"], results, export_start_idx, export_end_idx)
    wb.save(out_file)


def export_result3_template(template_file, out_file, results, export_start_idx, export_end_idx,
                             price_avg, n_days):
    """问题3专用：多一张"调整购电量"sheet（其余三张与result2同格式）。
    results[i] 需含：g_plan（0点计划）、g_final（最终执行/调整后）、
    cv/dv/socv（最终充放电与储电量轨迹）、deficit_e（实时缺口）。"""
    from openpyxl import load_workbook
    wb = load_workbook(template_file)
    _fill_purchase_sheet(wb["计划购电量"], results, "g_plan", export_start_idx, export_end_idx, price_avg, n_days)
    _fill_purchase_sheet(wb["调整购电量"], results, "g_final", export_start_idx, export_end_idx, price_avg, n_days)
    _fill_charge_sheet(wb["充放电量"], results, export_start_idx, export_end_idx)
    _fill_emergency_sheet(wb["紧急购电量"], results, export_start_idx, export_end_idx)
    wb.save(out_file)


# ======================= 论文素材：汇总报告 + 代表日画图 =======================
def build_summary_report(header_lines, results, dates_str, export_start_idx, export_end_idx,
                          price_avg, rep_dates=REP_DATES):
    exp_slice = results[export_start_idx:export_end_idx + 1]
    f_load_mat = np.array([r["f_load_kw"] for r in exp_slice])
    a_load_mat = np.array([r["actual_load_kw"] for r in exp_slice])
    f_pv_mat = np.array([r["f_pv_kw"] for r in exp_slice])
    a_pv_mat = np.array([r["actual_pv_kw"] for r in exp_slice])

    load_mae, load_rmse, load_mape = error_metrics(f_load_mat, a_load_mat)
    pv_mae, pv_rmse, pv_mape = error_metrics(f_pv_mat, a_pv_mat, mape_thresh=50.0)

    total_plan_cost = sum(r["plan_cost"] for r in exp_slice)
    total_emerg_cost = sum(r["emerg_cost"] for r in exp_slice)
    emerg_days = sum(1 for r in exp_slice if r["emerg_cost"] > 1e-6)
    emerg_blocks = sum(int(np.sum(r["deficit_e"] > 1e-6)) for r in exp_slice)

    metrics = dict(load_mae=load_mae, load_rmse=load_rmse, load_mape=load_mape,
                   pv_mae=pv_mae, pv_rmse=pv_rmse, pv_mape=pv_mape,
                   total_plan_cost=total_plan_cost, total_emerg_cost=total_emerg_cost,
                   total_cost=total_plan_cost + total_emerg_cost,
                   emerg_days=emerg_days, n_days=len(exp_slice), emerg_blocks=emerg_blocks)

    lines = []
    lines.append("=" * 70)
    lines.append("问题2  结果汇总")
    lines.append("=" * 70)
    lines.extend(header_lines)
    lines.append("")
    lines.append(f"预测误差（{EXPORT_START}~{EXPORT_END}）")
    lines.append("-" * 70)
    lines.append(f"  负荷  MAE={load_mae:8.2f} kW   RMSE={load_rmse:8.2f} kW   MAPE={load_mape:6.2f}%")
    lines.append(f"  光伏  MAE={pv_mae:8.2f} kW   RMSE={pv_rmse:8.2f} kW   MAPE(出力>50kW)={pv_mape:6.2f}%")
    lines.append("")
    lines.append(f"全年费用（{EXPORT_START}~{EXPORT_END}，{len(exp_slice)}天）")
    lines.append("-" * 70)
    lines.append(f"  计划购电费合计 = {total_plan_cost:12.2f} 元")
    lines.append(f"  紧急购电费合计 = {total_emerg_cost:12.2f} 元")
    lines.append(f"  总费用         = {total_plan_cost + total_emerg_cost:12.2f} 元")
    lines.append(f"  发生紧急购电天数 = {emerg_days} / {len(exp_slice)}")
    lines.append(f"  发生紧急购电时段数 = {emerg_blocks}")
    lines.append("")

    for rd in rep_dates:
        idx = int(np.where(dates_str == rd)[0][0])
        if idx >= len(results):
            continue
        r = results[idx]
        gv, cv, dv, socv, deficit_e = r["gv"], r["cv"], r["dv"], r["socv"], r["deficit_e"]
        lines.append("=" * 70)
        lines.append(f"代表日 {rd}")
        lines.append("=" * 70)
        lines.append("表1  微网在指定时间段的购电量 (kWh)")
        for name, t in TBL1:
            lines.append(f"  {name:<14s} {gv[t]:12.4f}")
        lines.append(f"  {'全天购电量 (kWh)':<14s} {gv.sum():12.4f}")
        lines.append(f"  {'全天购电费 (元)':<14s} {float(price_avg @ gv):12.2f}")
        lines.append("")
        lines.append("表2  储能设备在指定时间段的充放电量 (kWh)")
        lines.append(f"  {'时间段':<14s} {'充电量':>12s} {'放电量':>12s}")
        for name, a, b in TBL2:
            lines.append(f"  {name:<14s} {cv[a:b].sum():12.4f} {dv[a:b].sum():12.4f}")
        lines.append(f"  0:00 储电量 (kWh)  = {socv[0]:.4f}")
        lines.append(f"  24:00 储电量 (kWh) = {socv[-1]:.4f}")
        lines.append("")
        lines.append("表3  紧急购电")
        segs = merge_deficit_segments(deficit_e)
        if not segs:
            lines.append("  (无紧急购电)")
        else:
            for s, e, val in segs:
                seg_label = time_range_str(s)[:5] + "-" + time_range_str(e - 1)[6:]
                lines.append(f"  {seg_label:<14s} {val:12.4f}")
        lines.append("")

    return "\n".join(lines), metrics


def plot_representative_days(results, dates_str, price_raw, out_dir, title_prefix="问题2",
                              rep_dates=REP_DATES, file_prefix="plot"):
    import matplotlib.pyplot as plt
    try:
        plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei", "DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
    except Exception:
        pass

    h = np.arange(T) * DT
    saved = []
    for rd in rep_dates:
        idx = int(np.where(dates_str == rd)[0][0])
        if idx >= len(results):
            continue
        r = results[idx]
        gv, cv, dv, socv = r["gv"], r["cv"], r["dv"], r["socv"]
        fig, ax = plt.subplots(3, 1, figsize=(11, 9), sharex=True)
        ax[0].plot(h, price_raw, color="tab:red")
        ax[0].set_ylabel("电价 (元/kWh)")
        ax[0].set_title(f"{title_prefix}：{rd} 电价 / 功率平衡 / 储能电量")
        ax[1].plot(h, r["actual_load_kw"], label="实际负荷", color="k")
        ax[1].plot(h, r["actual_pv_kw"], label="实际光伏", color="tab:orange")
        ax[1].plot(h, r["f_load_kw"], label="预测负荷", color="k", ls="--", alpha=0.6)
        ax[1].plot(h, r["f_pv_kw"], label="预测光伏", color="tab:orange", ls="--", alpha=0.6)
        ax[1].plot(h, gv / DT, label="计划购电功率", color="tab:blue")
        ax[1].plot(h, (dv - cv) / DT, label="储能净放电功率", color="tab:green")
        ax[1].legend(ncol=3, fontsize=7)
        ax[1].set_ylabel("功率 (kW)")
        ax[2].plot(np.arange(T + 1) * DT, socv, color="tab:purple")
        ax[2].axhline(SOC_MIN, ls="--", c="gray")
        ax[2].axhline(SOC_MAX, ls="--", c="gray")
        ax[2].set_ylabel("储电量 (kWh)")
        ax[2].set_xlabel("时刻 (h)")
        fig.tight_layout()
        out_png = Path(out_dir) / f"{file_prefix}_{rd.replace('-', '')}.png"
        fig.savefig(out_png, dpi=150)
        plt.close(fig)
        saved.append(out_png)
    return saved
